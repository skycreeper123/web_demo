import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
OUT_TECH = BASE_DIR / "video_batch_technical_doc.docx"
OUT_USER = BASE_DIR / "video_batch_user_guide.docx"
SKILL_SCRIPTS_DIR = Path(
    "C:/Users/but48/.codex/plugins/cache/openai-primary-runtime/documents/26.813.12317/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS_DIR))

from table_geometry import apply_table_geometry, exact_column_widths

FONT_LATIN = "Calibri"
FONT_EAST_ASIA = "Microsoft YaHei"
TITLE_COLOR = RGBColor(31, 77, 120)
HEADING_COLOR = RGBColor(46, 116, 181)
BODY_COLOR = RGBColor(0, 0, 0)
MUTED_COLOR = RGBColor(85, 85, 85)
TABLE_HEADER_FILL = "E8EEF5"
TABLE_BORDER_COLOR = "D9E2F2"

VIDEO_EXTS = ".mp4, .mov, .mkv, .avi, .webm, .m4v"


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT_LATIN
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(paragraph, before=0, after=6, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    paragraph.alignment = align


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=TABLE_BORDER_COLOR, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tc_borders.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.clear()


def add_title(document, text, subtitle=None):
    p = document.add_paragraph()
    set_paragraph_format(p, before=0, after=3, line=1.0)
    run = p.add_run(text)
    set_run_font(run, size=24, bold=True, color=TITLE_COLOR)

    if subtitle:
        sp = document.add_paragraph()
        set_paragraph_format(sp, before=0, after=10, line=1.0)
        run = sp.add_run(subtitle)
        set_run_font(run, size=10.5, color=MUTED_COLOR)


def style_body_style(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_COLOR
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, HEADING_COLOR, 18, 10),
        ("Heading 2", 13, HEADING_COLOR, 14, 7),
        ("Heading 3", 12, TITLE_COLOR, 10, 5),
    ]:
        style = document.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.25


def set_section_geometry(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    margin = Inches(1)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def make_bullet(document, text, level=0):
    p = document.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    set_paragraph_format(p, before=0, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def make_number(document, text, style_name="List Number"):
    p = document.add_paragraph(style=style_name)
    set_paragraph_format(p, before=0, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_heading(document, text, level):
    p = document.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=False, color={1: HEADING_COLOR, 2: HEADING_COLOR, 3: TITLE_COLOR}[level])
    return p


def add_label_detail_table(document, rows, widths):
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
        for idx, cell in enumerate(row):
            for paragraph in cell.paragraphs:
                set_paragraph_format(paragraph, before=0, after=0, line=1.1)
                for run in paragraph.runs:
                    set_run_font(run, size=10.5)
            if idx == 0:
                set_cell_shading(cell, "F4F6F9")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
        set_cell_border(row[0])
        set_cell_border(row[1])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    return table


def add_script_table(document, rows, widths):
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["脚本", "方式", "输入", "输出", "核心行为"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        set_cell_shading(cell, TABLE_HEADER_FILL)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(p, before=0, after=0, line=1.0)
            for run in p.runs:
                set_run_font(run, size=10.5, bold=True)
        set_cell_border(cell)
    for row_data in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value
            for p in row[idx].paragraphs:
                set_paragraph_format(p, before=0, after=0, line=1.1)
                if idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=10)
            set_cell_border(row[idx])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    return table


def add_three_col_table(document, rows, widths):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["场景", "推荐脚本", "备注"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        set_cell_shading(cell, TABLE_HEADER_FILL)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(p, before=0, after=0, line=1.0)
            for run in p.runs:
                set_run_font(run, size=10.5, bold=True)
        set_cell_border(cell)
    for row_data in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data):
            row[idx].text = value
            for p in row[idx].paragraphs:
                set_paragraph_format(p, before=0, after=0, line=1.1)
                for run in p.runs:
                    set_run_font(run, size=10)
            set_cell_border(row[idx])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    return table


def add_two_col_table(document, rows, widths):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], "F4F6F9")
        for idx, cell in enumerate(cells):
            for p in cell.paragraphs:
                set_paragraph_format(p, before=0, after=0, line=1.1)
                for run in p.runs:
                    set_run_font(run, size=10.5)
                if idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_cell_border(cell)
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    return table


def add_note(document, text):
    p = document.add_paragraph()
    set_paragraph_format(p, before=3, after=6, line=1.2)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=MUTED_COLOR, italic=True)


def build_technical_doc(path: Path):
    doc = Document()
    set_section_geometry(doc)
    style_body_style(doc)
    add_title(doc, "视频批处理脚本技术文档", "脚本功能梳理、实现方式说明与输入输出约定")

    add_heading(doc, "1. 概述", 1)
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=6, line=1.25)
    r = p.add_run("本目录脚本分为两类：一类负责对单个视频批量保留前段或后段，另一类负责把两个文件夹中的视频按顺序配对合并。")
    set_run_font(r, size=11)

    for bullet in [
        "默认输入目录为脚本同级的 `input_videos`，合并脚本使用 `input_folder_a` 和 `input_folder_b`。",
        "输出目录由脚本内部固定命名并自动创建，例如 `output_first_3s`、`output_tail_30pct`、`output_merged_pairs`。",
        "所有脚本都会按文件扩展名过滤视频，支持 `.mp4`、`.mov`、`.mkv`、`.avi`、`.webm`、`.m4v`。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "2. 依赖与处理方式", 1)
    add_label_detail_table(
        doc,
        [
            ("Python", "用于运行所有脚本。"),
            ("OpenCV", "除 `video_batch_half_keep.py` 外，其余视频处理脚本都依赖 `cv2`。"),
            ("ffmpeg / ffprobe", "`video_batch_half_keep.py` 需要 `ffmpeg` 和 `ffprobe` 在 PATH 中。"),
            ("输出编码", "OpenCV 版本统一写出为 `mp4v` 编码的 MP4 文件。"),
        ],
        exact_column_widths([2200, 7160]),
    )

    add_heading(doc, "3. 脚本清单", 1)
    add_script_table(
        doc,
        [
            ("video_batch_half_keep.py", "ffmpeg + ffprobe", "input_videos", "output/trimmed, output/last_frames", "保留前 50% 时长，视频使用 `-c copy` 快速裁切，并在中点附近抽取预览帧。"),
            ("video_batch_half_keep_opencv.py", "OpenCV 按帧写入", "input_videos", "output_opencv/trimmed, output_opencv/last_frames", "保留前 50% 帧并重新编码，预览图为保留段最后一帧。"),
            ("video_batch_keep_30pct.py", "OpenCV 按帧写入", "input_videos", "output_30pct/trimmed, output_30pct/last_frames", "保留前 30% 帧。"),
            ("video_batch_keep_70pct.py", "OpenCV 按帧写入", "input_videos", "output_70pct/trimmed, output_70pct/last_frames", "保留前 70% 帧。"),
            ("video_batch_keep_first_3s.py", "OpenCV 按时长写入", "input_videos", "output_first_3s/trimmed, output_first_3s/last_frames", "保留前 3 秒。"),
            ("video_batch_keep_first_5s.py", "OpenCV 按时长写入", "input_videos", "output_first_5s/trimmed, output_first_5s/last_frames", "保留前 5 秒。"),
            ("video_batch_keep_first_7s.py", "OpenCV 按时长写入", "input_videos", "output_first_7s/trimmed, output_first_7s/last_frames", "保留前 7 秒。"),
            ("video_batch_keep_last_3s.py", "OpenCV 反向定位", "input_videos", "output_last_3s/trimmed, output_last_3s/first_frames", "保留后 3 秒，预览图为截取段起始帧。"),
            ("video_batch_keep_last_5s.py", "OpenCV 反向定位", "input_videos", "output_last_5s/trimmed, output_last_5s/first_frames", "保留后 5 秒，预览图为截取段起始帧。"),
            ("video_batch_keep_last_7s.py", "OpenCV 反向定位", "input_videos", "output_last_7s/trimmed, output_last_7s/first_frames", "保留后 7 秒，预览图为截取段起始帧。"),
            ("video_batch_keep_tail_30pct.py", "OpenCV 反向定位", "input_videos", "output_tail_30pct/trimmed, output_tail_30pct/first_frames", "保留后 30% 帧。"),
            ("video_batch_keep_tail_50pct.py", "OpenCV 反向定位", "input_videos", "output_tail_50pct/trimmed, output_tail_50pct/first_frames", "保留后 50% 帧。"),
            ("video_batch_keep_tail_70pct.py", "OpenCV 反向定位", "input_videos", "output_tail_70pct/trimmed, output_tail_70pct/first_frames", "保留后 70% 帧。"),
            ("merge_two_folders_pairwise.py", "OpenCV 双文件夹合并", "input_folder_a, input_folder_b", "output_merged_pairs", "按排序后一一配对合并，第二个文件夹的视频会被缩放到第一个文件夹的视频尺寸。"),
        ],
        exact_column_widths([1880, 1200, 1500, 1840, 2940]),
    )

    add_heading(doc, "4. 实现特征", 1)
    for bullet in [
        "前半段脚本中，`ffmpeg` 版本更快，但切点会受关键帧影响；OpenCV 版本更稳定，代价是会重新编码。",
        "后半段脚本通过先定位到起始帧，再顺序写出剩余帧，因此预览图保存的是保留区间的第一帧。",
        "所有脚本都采用固定输出文件名，重复运行时会覆盖同名结果。",
        "合并脚本按文件名排序后逐对处理，只合并两侧共同拥有的前 `min(len(A), len(B))` 对文件。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "5. 适用场景建议", 1)
    add_label_detail_table(
        doc,
        [
            ("追求速度", "优先使用 `video_batch_half_keep.py`。"),
            ("追求帧级稳定", "优先使用 OpenCV 版本。"),
            ("固定秒数", "使用 `*_first_3s.py`、`*_first_5s.py`、`*_first_7s.py` 或对应 `last` 脚本。"),
            ("固定比例", "使用 `*_30pct.py`、`*_50pct.py`、`*_70pct.py` 或对应 `tail` 脚本。"),
            ("双文件夹拼接", "使用 `merge_two_folders_pairwise.py`。"),
        ],
        exact_column_widths([2000, 7360]),
    )

    add_note(doc, "说明：本技术文档只整理当前目录已有脚本的实际行为，不额外假设命令行参数或可配置项。")
    doc.save(path)


def build_user_doc(path: Path):
    doc = Document()
    set_section_geometry(doc)
    style_body_style(doc)
    add_title(doc, "视频批处理脚本使用文档", "按场景快速选择脚本并完成批量裁切或合并")

    add_heading(doc, "1. 使用前准备", 1)
    for bullet in [
        "确保已安装 Python，并且 `opencv-python` 可正常导入。",
        "如果要使用 `video_batch_half_keep.py`，还需要提前安装 `ffmpeg` 和 `ffprobe`，并加入 PATH。",
        "把待处理视频放进对应输入文件夹：单文件夹脚本使用 `input_videos`，合并脚本使用 `input_folder_a` 和 `input_folder_b`。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "2. 文件夹约定", 1)
    add_two_col_table(
        doc,
        [
            ("input_videos", "放入需要批量裁切的原始视频。大多数脚本都会自动读取这里的所有视频文件。"),
            ("input_folder_a / input_folder_b", "放入需要按顺序合并的两组视频。两个文件夹的视频数量可以不同，但只会合并成对部分。"),
            ("output_*", "脚本会自动创建输出目录；若同名结果已存在，通常会被覆盖。"),
        ],
        exact_column_widths([2300, 7060]),
    )

    add_heading(doc, "3. 先选脚本", 1)
    add_three_col_table(
        doc,
        [
            ("想要最快裁前半段", "video_batch_half_keep.py", "保留前 50% 时长，速度最快。"),
            ("想要更稳的帧级裁切", "video_batch_half_keep_opencv.py", "OpenCV 按帧输出，适合需要稳定结果的情况。"),
            ("保留前几秒", "video_batch_keep_first_3s.py / 5s / 7s.py", "分别保留前 3、5、7 秒。"),
            ("保留后几秒", "video_batch_keep_last_3s.py / 5s / 7s.py", "分别保留后 3、5、7 秒。"),
            ("保留前几个百分比", "video_batch_keep_30pct.py / 70pct.py", "分别保留前 30% 或 70%。"),
            ("保留后几个百分比", "video_batch_keep_tail_30pct.py / 50pct.py / 70pct.py", "分别保留后 30%、50% 或 70%。"),
            ("合并两个文件夹", "merge_two_folders_pairwise.py", "按排序后一一配对合并。"),
        ],
        exact_column_widths([2100, 3000, 4260]),
    )

    add_heading(doc, "4. 常用操作", 1)
    add_heading(doc, "4.1 保留前 N 秒 / 前 N%", 2)
    for bullet in [
        "把原始视频放进 `input_videos`。",
        "选择对应脚本并运行。",
        "结果会出现在脚本自动创建的 `output_* / trimmed` 目录中，预览图在对应的 `last_frames` 目录中。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "4.2 保留后 N 秒 / 后 N%", 2)
    for bullet in [
        "把原始视频放进 `input_videos`。",
        "选择对应 `last` 或 `tail` 脚本运行。",
        "结果视频在 `trimmed` 目录，预览图通常是保留区间的第一帧。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "4.3 两个文件夹合并", 2)
    for bullet in [
        "把第一组视频放进 `input_folder_a`，第二组视频放进 `input_folder_b`。",
        "确保文件名排序顺序就是你想要的配对顺序。",
        "运行 `merge_two_folders_pairwise.py` 后，输出会写到 `output_merged_pairs`。",
    ]:
        make_bullet(doc, bullet)

    add_heading(doc, "5. 输出命名", 1)
    add_label_detail_table(
        doc,
        [
            ("裁切视频", "统一输出为 `*_first_*`、`*_last_*` 或 `*_tail_*` 命名的 MP4。"),
            ("预览图片", "统一输出为 JPG，文件名会标明是保留段的最后一帧或第一帧。"),
            ("合并结果", "命名格式为 `原始A文件名__原始B文件名.mp4`。"),
        ],
        exact_column_widths([2000, 7360]),
    )

    add_heading(doc, "6. 常见问题", 1)
    for bullet in [
        "提示 `Cannot open`，通常是视频编码不被 OpenCV 支持，或文件已损坏。",
        "提示找不到 `ffmpeg` / `ffprobe`，先把它们安装并加入 PATH。",
        "合并脚本如果两边数量不同，只会处理较短一侧能配对上的部分。",
        "重复运行同一脚本会覆盖同名输出，建议先确认输出目录里没有需要保留的旧文件。",
        "如果对裁切边界要求非常严格，优先用 OpenCV 版本而不是 `ffmpeg -c copy` 版本。",
    ]:
        make_bullet(doc, bullet)

    add_note(doc, "提示：这套脚本更适合“批量固定规则处理”。如果你后面想要加参数化命令行，我也可以继续帮你改成统一入口。")
    doc.save(path)


def main():
    build_technical_doc(OUT_TECH)
    build_user_doc(OUT_USER)
    print(f"Created: {OUT_TECH}")
    print(f"Created: {OUT_USER}")


if __name__ == "__main__":
    main()
